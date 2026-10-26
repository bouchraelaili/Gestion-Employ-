import io
import psycopg2
from .base_service import BaseService


class ExportService(BaseService):

    @staticmethod
    def get_employees_data():
        conn = BaseService.get_db()
        cur  = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cur.execute("""
            SELECT e.id, e.nom, e.prenom, e.age, e.sexe,
                   c.nom AS categorie, p.nom AS poste,
                   e.date_embauche,
                   DATE_PART('year', AGE(e.date_embauche))::INT AS anciennete,
                   e.salaire, e.absences, e.score_performance,
                   CASE WHEN e.risque_depart THEN 'Oui' ELSE 'Non' END AS risque_depart
            FROM employees e
            JOIN categories c ON e.category_id = c.id
            JOIN postes     p ON e.poste_id     = p.id
            ORDER BY e.id
        """)
        rows = cur.fetchall()
        cur.close(); conn.close()
        return rows

    #  EXCEL
    @staticmethod
    def generate_excel():
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        except ImportError:
            raise ImportError("pip install openpyxl")

        from datetime import date

        rows = ExportService.get_employees_data()
        wb   = openpyxl.Workbook()
        ws   = wb.active
        ws.title = "Rapport RH"

        DARK  = "1A1A2E"; PINK = "E84393"; LIGHT = "F8F9FC"
        BORDER= "E8EAF0"; WHITE= "FFFFFF"; TEXT  = "1A1A2E"
        RED   = "991B1B"; RED_BG="FEE2E2"; GREEN = "0F6E56"; GREEN_BG="E0F7F0"

        widths = [5, 14, 14, 6, 6, 22, 22, 10, 14, 14, 10, 8, 10]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

        thin   = Side(style="thin", color=BORDER)
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws.merge_cells("A1:M1")
        ws["A1"] = "RAPPORT RH — SYSTEME DE GESTION DES RESSOURCES HUMAINES"
        ws["A1"].font      = Font(name="Calibri", bold=True, size=14, color=WHITE)
        ws["A1"].fill      = PatternFill("solid", fgColor=DARK)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:M2")
        ws["A2"] = f"Genere le {date.today().strftime('%d/%m/%Y')}  |  Total : {len(rows)} employes"
        ws["A2"].font      = Font(name="Calibri", size=10, color="FF6B9D")
        ws["A2"].fill      = PatternFill("solid", fgColor=DARK)
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22
        ws.row_dimensions[3].height = 6

        headers = ["ID","Nom","Prenom","Age","Sexe","Categorie","Poste",
                   "Anciennete","Salaire (DA)","Date embauche","Absences","Score","Risque"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=h)
            cell.font      = Font(name="Calibri", bold=True, size=10, color=WHITE)
            cell.fill      = PatternFill("solid", fgColor=PINK)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = border
        ws.row_dimensions[4].height = 26

        for r_idx, row in enumerate(rows, 5):
            bg     = WHITE if r_idx % 2 == 0 else LIGHT
            risque = row["risque_depart"] == "Oui"
            values = [
                row["id"], row["nom"], row["prenom"], row["age"], row["sexe"],
                row["categorie"], row["poste"], f'{row["anciennete"]} ans',
                f'{float(row["salaire"]):,.0f}', str(row["date_embauche"])[:10],
                row["absences"], row["score_performance"], row["risque_depart"]
            ]
            for c_idx, val in enumerate(values, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font      = Font(name="Calibri", size=10, color=TEXT)
                cell.fill      = PatternFill("solid", fgColor=bg)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border    = border
                if c_idx == 13:
                    cell.font = Font(name="Calibri", bold=True, size=10,
                                     color=RED if risque else GREEN)
                    cell.fill = PatternFill("solid", fgColor=RED_BG if risque else GREEN_BG)
                if c_idx == 12 and isinstance(val, int) and val < 55:
                    cell.font = Font(name="Calibri", bold=True, size=10, color=PINK)
                if c_idx == 11 and isinstance(val, int) and val > 12:
                    cell.font = Font(name="Calibri", bold=True, size=10, color=RED)
            ws.row_dimensions[r_idx].height = 20

        ws.freeze_panes = "A5"
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    #  PDF
    @staticmethod
    def generate_pdf():
        try:
            from reportlab.lib.pagesizes import landscape, A3
            from reportlab.lib           import colors
            from reportlab.lib.units     import mm
            from reportlab.lib.styles    import ParagraphStyle
            from reportlab.lib.enums     import TA_CENTER, TA_LEFT, TA_RIGHT
            from reportlab.platypus      import (
                SimpleDocTemplate, Table, TableStyle,
                Paragraph, Spacer, HRFlowable
            )
        except ImportError:
            raise ImportError("pip install reportlab")

        from datetime import date as dt

        rows   = ExportService.get_employees_data()
        output = io.BytesIO()

        PAGE_W = landscape(A3)[0]

        doc = SimpleDocTemplate(
            output, pagesize=landscape(A3),
            leftMargin=15*mm, rightMargin=15*mm,
            topMargin=12*mm,  bottomMargin=12*mm,
        )

        C = {
            "dark":       colors.HexColor("#1A1A2E"),
            "dark2":      colors.HexColor("#16213E"),
            "pink":       colors.HexColor("#E84393"),
            "pink_mid":   colors.HexColor("#FCE4F0"),
            "pink_light": colors.HexColor("#FDF2F8"),
            "pink2":      colors.HexColor("#FF6B9D"),
            "row_alt":    colors.HexColor("#F5F6FA"),
            "border":     colors.HexColor("#E8EAF0"),
            "green":      colors.HexColor("#0F6E56"),
            "green_bg":   colors.HexColor("#E0F7F0"),
            "green_brd":  colors.HexColor("#86EFAC"),
            "red":        colors.HexColor("#991B1B"),
            "red_bg":     colors.HexColor("#FEE2E2"),
            "red_brd":    colors.HexColor("#FCA5A5"),
            "amber":      colors.HexColor("#D97706"),
            "amber_bg":   colors.HexColor("#FEF3C7"),
            "amber_brd":  colors.HexColor("#FCD34D"),
            "white":      colors.white,
            "gray":       colors.HexColor("#8B8FA8"),
            "gray_light": colors.HexColor("#F0F2F8"),
            "text":       colors.HexColor("#1A1A2E"),
        }

        USABLE_W = PAGE_W - 30*mm

        def ps(name, **kw):
            defaults = dict(fontName="Helvetica", fontSize=9,
                            textColor=C["text"], alignment=TA_CENTER, leading=12)
            defaults.update(kw)
            return ParagraphStyle(name, **defaults)

        elements = []

        #  HEADER 
        total_risque = sum(1 for r in rows if r["risque_depart"] == "Oui")
        stable       = len(rows) - total_risque
        avg_score    = round(sum(r["score_performance"] for r in rows)/len(rows)) if rows else 0
        avg_abs      = round(sum(r["absences"] for r in rows)/len(rows), 1) if rows else 0

        banner = Table([
            [Paragraph(
                "RAPPORT RH — SYSTEME DE GESTION DES RESSOURCES HUMAINES",
                ps("tt", fontName="Helvetica-Bold", fontSize=16,
                   textColor=C["white"], leading=20)
            )],
            [Paragraph(
                f"Genere le {dt.today().strftime('%d/%m/%Y')}"
                f"&nbsp;&nbsp;|&nbsp;&nbsp;{len(rows)} employes"
                f"&nbsp;&nbsp;|&nbsp;&nbsp;{total_risque} a risque",
                ps("dt", fontSize=10, textColor=C["pink2"], leading=14)
            )],
        ], colWidths=[USABLE_W])
        banner.setStyle(TableStyle([
            ("BACKGROUND",    (0,0),(0,0), C["dark"]),
            ("BACKGROUND",    (0,1),(0,1), C["dark2"]),
            ("TOPPADDING",    (0,0),(0,0), 16),
            ("BOTTOMPADDING", (0,0),(0,0), 8),
            ("TOPPADDING",    (0,1),(0,1), 6),
            ("BOTTOMPADDING", (0,1),(0,1), 14),
            ("LEFTPADDING",   (0,0),(-1,-1), 20),
            ("RIGHTPADDING",  (0,0),(-1,-1), 20),
        ]))
        elements.append(banner)
        elements.append(Spacer(1, 10))

        #  KPI CARDS 
        kpi_items = [
            ("Total employes",  str(len(rows)),    C["pink"],  C["pink_light"], C["pink_mid"]),
            ("A risque",        str(total_risque), C["red"],   C["red_bg"],     C["red_brd"]),
            ("Stables",         str(stable),       C["green"], C["green_bg"],   C["green_brd"]),
            ("Score moyen",     f"{avg_score}/100",C["dark"],  C["gray_light"], C["border"]),
            ("Absences moy.",   f"{avg_abs} j",    C["amber"], C["amber_bg"],   C["amber_brd"]),
        ]
        kpi_w = USABLE_W / 5

        kpi_cells = []
        for label, val, fg, bg, brd in kpi_items:
            t = Table([
                [Paragraph(val,   ps(f"kv{label}", fontName="Helvetica-Bold",
                                     fontSize=26, textColor=fg, leading=30))],
                [Paragraph(label, ps(f"kl{label}", fontSize=9,
                                     textColor=C["gray"], leading=12))],
            ], colWidths=[kpi_w - 3*mm])
            t.setStyle(TableStyle([
                ("BACKGROUND",    (0,0),(-1,-1), bg),
                ("TOPPADDING",    (0,0),(-1,-1), 12),
                ("BOTTOMPADDING", (0,0),(-1,-1), 12),
                ("LEFTPADDING",   (0,0),(-1,-1), 10),
                ("RIGHTPADDING",  (0,0),(-1,-1), 10),
                ("BOX",           (0,0),(-1,-1), 1.5, brd),
            ]))
            kpi_cells.append(t)

        kpi_row = Table([kpi_cells], colWidths=[kpi_w]*5)
        kpi_row.setStyle(TableStyle([
            ("LEFTPADDING",  (0,0),(-1,-1), 2),
            ("RIGHTPADDING", (0,0),(-1,-1), 2),
            ("TOPPADDING",   (0,0),(-1,-1), 0),
            ("BOTTOMPADDING",(0,0),(-1,-1), 0),
        ]))
        elements.append(kpi_row)
        elements.append(Spacer(1, 12))

        #  TABLEAU 
        def h(txt):
            return Paragraph(f"<b>{txt}</b>",
                             ps(f"h{txt}", fontName="Helvetica-Bold",
                                fontSize=9, textColor=C["white"]))

        def c(txt, color=None, bold=False, align=TA_CENTER):
            return Paragraph(
                f"<b>{txt}</b>" if bold else str(txt),
                ps(f"c{txt}", fontName="Helvetica-Bold" if bold else "Helvetica",
                   fontSize=9, textColor=color or C["text"],
                   alignment=align, leading=11)
            )

       
        cw_mm = [10, 25, 25, 10, 10, 50, 52, 16, 32, 26, 14, 14, 18]
        total_mm = sum(cw_mm)
        scale    = (USABLE_W / mm) / total_mm
        cw       = [w * scale * mm for w in cw_mm]

        data = [[
            h("ID"), h("Nom"), h("Prenom"), h("Age"), h("Sexe"),
            h("Categorie"), h("Poste"), h("Ancien."),
            h("Salaire DA"), h("Embauche"),
            h("Abs."), h("Score"), h("Risque"),
        ]]

        risk_rows = []
        for i, row in enumerate(rows, 1):
            risque   = row["risque_depart"] == "Oui"
            score    = row["score_performance"]
            absences = row["absences"]
            sc  = C["red"] if score < 50 else (C["amber"] if score < 65 else C["green"])
            abc = C["red"] if absences > 15 else (C["amber"] if absences > 10 else C["green"])

            data.append([
                c(row["id"]),
                c(row["nom"],    align=TA_LEFT),
                c(row["prenom"], align=TA_LEFT),
                c(row["age"]),
                c(row["sexe"]),
                c(row["categorie"], align=TA_LEFT),
                c(row["poste"],     align=TA_LEFT),
                c(f'{row["anciennete"]}'),
                c(f'{float(row["salaire"]):,.0f}', align=TA_RIGHT),
                c(str(row["date_embauche"])[:10]),
                c(absences,  color=abc,  bold=absences > 10),
                c(score,     color=sc,   bold=True),
                c("Oui" if risque else "Non",
                  color=C["red"] if risque else C["green"], bold=True),
            ])
            if risque:
                risk_rows.append(i)

        tbl = Table(data, colWidths=cw, repeatRows=1)
        tbl_style = [
            ("BACKGROUND",    (0,0),  (-1,0),  C["dark"]),
            ("LINEBELOW",     (0,0),  (-1,0),  2, C["pink"]),
            ("TOPPADDING",    (0,0),  (-1,-1), 7),
            ("BOTTOMPADDING", (0,0),  (-1,-1), 7),
            ("LEFTPADDING",   (0,0),  (-1,-1), 6),
            ("RIGHTPADDING",  (0,0),  (-1,-1), 6),
            ("GRID",          (0,0),  (-1,-1), 0.4, C["border"]),
            ("VALIGN",        (0,0),  (-1,-1), "MIDDLE"),
            ("ROWBACKGROUNDS",(0,1),  (-1,-1), [C["white"], C["row_alt"]]),
        ]
        for r in risk_rows:
            tbl_style += [
                ("BACKGROUND", (12,r),(12,r), C["red_bg"]),
                ("BACKGROUND", (10,r),(10,r), C["red_bg"]),
            ]
        tbl.setStyle(TableStyle(tbl_style))
        elements.append(tbl)

        #  FOOTER 
        elements.append(Spacer(1, 12))
        elements.append(HRFlowable(
            width="100%", thickness=1, color=C["pink"], spaceAfter=5
        ))
        elements.append(Paragraph(
            f"RH System &nbsp;|&nbsp; Rapport confidentiel"
            f" &nbsp;|&nbsp; {dt.today().strftime('%d/%m/%Y')}"
            f" &nbsp;|&nbsp; Document genere automatiquement",
            ps("footer", fontSize=8, textColor=C["gray"])
        ))

        doc.build(elements)
        output.seek(0)
        return output